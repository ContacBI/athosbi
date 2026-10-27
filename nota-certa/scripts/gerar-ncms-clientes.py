"""Lista de NCMs/NBS vendidos pelos clientes (resumo fiscal da Domínio) com o
enquadramento na LC 214 — para escolher os pilotos e começar a matriz fiscal.

Entrada: um JSON com os clientes. O resumo vem da tabela dominio_fiscal, que a
Central manda só lendo a Domínio. Formato:
  { "sincronizado": "dd/mm/aaaa",
    "clientes": [ { "nome", "aba", "cnpj", "inicio", "fim", "meses",
                    "observacoes": [..],
                    "vendas": [ { "ncm", "servico", "descricao", "produtos", "valor", "cfops", "pis_cofins" } ],
                    "leituras": { "<código>": ["<classificação>", "<leitura inicial>"] } } ] }

Uso: python3 gerar-ncms-clientes.py entrada.json "saida.xlsx"

ATENÇÃO: a entrada e a planilha gerada têm dados reais de clientes e NÃO vão
para o repositório. Este script é só o gerador.
"""

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DADOS = Path(__file__).resolve().parent.parent / "dados"
NAVY, MUTED, INPUT, RED, FONT = "14202B", "5B6670", "FFF4D6", "B42318", "Arial"
CATEGORIAS = [
    ("Zero (qualquer comprador)", "E6F4EC"),
    ("60% (zero p/ órgão público/SUS)", "EAF1FB"),
    ("Revisar (provável alíquota cheia)", "FFF4D6"),
    ("Alíquota cheia", "F3F5F7"),
]
thin = Side(style="thin", color="D5DBE0")
BORDER = Border(top=thin, bottom=thin, left=thin, right=thin)


def font(**kw):
    return Font(name=FONT, size=kw.pop("size", 10), **kw)


def fmt_codigo(codigo, servico):
    d = "".join(ch for ch in codigo if ch.isdigit())
    if servico and len(d) == 9:
        return f"{d[0]}.{d[1:5]}.{d[5:7]}.{d[7:]}"
    if not servico and len(d) == 8:
        return f"{d[:4]}.{d[4:6]}.{d[6:]}"
    return codigo or "(sem código)"


def enquadramentos(anexos, codigo, servico):
    """Mesma regra da Calculadora: prefixo do anexo; exceção só exclui se não houver vínculo sem exceção."""
    d = "".join(ch for ch in codigo if ch.isdigit())
    if len(d) != (9 if servico else 8):
        return []
    por = {}
    for prefixo, cclass, item, excecoes in anexos["nbs" if servico else "ncm"]:
        if d.startswith(prefixo):
            por.setdefault(cclass, []).append((item, excecoes or []))
    out = []
    for cclass, vinculos in por.items():
        excluido = any(any(d.startswith(e) for e in ex) for _, ex in vinculos)
        limpo = any(not ex for _, ex in vinculos)
        if excluido and not limpo:
            continue
        t = anexos["tratamentos"][cclass]
        itens = sorted({i for i, _ in vinculos if i}, key=lambda i: int("".join(c for c in i.split("/")[1] if c.isdigit()) or 0))
        out.append({"cclass": cclass, **t, "itens": itens})
    return sorted(out, key=lambda e: (bool(e.get("condicao")), -e["reducao"]))


def rotulo(e):
    return f"{e['cclass']} — {e['titulo']} ({'alíquota zero' if e['reducao'] == 100 else 'redução de %d%%' % e['reducao']})"


def classificacao_auto(lista):
    principal = next((e for e in lista if not e.get("condicao")), None)
    if not lista:
        return "Alíquota cheia"
    if principal and principal["reducao"] == 100:
        return "Zero (qualquer comprador)"
    if principal and any(e.get("condicao") for e in lista):
        return "60% (zero p/ órgão público/SUS)"
    return "Revisar (provável alíquota cheia)"


def cabecalho(ws, linha, colunas, entrada=()):
    for c, (nome, largura) in enumerate(colunas, start=1):
        cel = ws.cell(row=linha, column=c, value=nome)
        cel.font = font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="8A5A00" if c in entrada else NAVY)
        cel.alignment = Alignment(wrap_text=True, vertical="center")
        cel.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = largura
    ws.row_dimensions[linha].height = 32


def aba_cliente(wb, anexos, cli, sincronizado):
    ws = wb.create_sheet(cli["aba"])
    periodo = f"{cli['inicio'][8:]}/{cli['inicio'][5:7]}/{cli['inicio'][:4]} a {cli['fim'][8:]}/{cli['fim'][5:7]}/{cli['fim'][:4]} ({cli['meses']} meses)"
    cnpj = cli["cnpj"]
    ws["A1"] = f"{cli['nome']} — vendas por NCM/NBS e enquadramento na LC 214"
    ws["A1"].font = font(size=14, bold=True, color=NAVY)
    ws["A2"] = (f"CNPJ {cnpj[:2]}.{cnpj[2:5]}.{cnpj[5:8]}/{cnpj[8:12]}-{cnpj[12:]} · período {periodo} · "
                f"fonte: resumo das notas fiscais da Domínio enviado pela Central (só leitura), sincronizado em {sincronizado}")
    ws["A2"].font = font(size=9, color=MUTED, italic=True)
    linha = 3
    for obs in cli["observacoes"]:
        cel = ws.cell(row=linha, column=1, value=f"• {obs}")
        cel.font = font(size=9, color="222222")
        cel.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=15)
        ws.row_dimensions[linha].height = 26
        linha += 1
    head = linha + 1
    colunas = [
        ("Código (NCM / NBS)", 14), ("Tipo", 9), ("Exemplo de produto (um dos vendidos)", 34), ("Produtos no código", 9),
        ("Vendas no período (R$)", 16), ("% das vendas", 9), ("CFOPs usados", 16), ("PIS/Cofins hoje", 9),
        ("Itens da LC 214 para este código", 46), ("Pelo NCM/NBS, a base sugere", 30), ("Regra com condição (só vale se…)", 28),
        ("Leitura inicial (a conferir)", 52), ("Resumo", 20), ("Decisão do escritório", 13), ("Observação", 30),
    ]
    cabecalho(ws, head, colunas, entrada=(14, 15))
    vendas = sorted(cli["vendas"], key=lambda v: -v["valor"])
    primeira = head + 1
    ultima = primeira + len(vendas) - 1
    for i, v in enumerate(vendas):
        r = primeira + i
        servico = bool(v["servico"]) or not v["ncm"]
        codigo = v["servico"] if servico else v["ncm"]
        lista = enquadramentos(anexos, codigo, servico)
        principal = next((e for e in lista if not e.get("condicao")), None)
        condicionado = next((e for e in lista if e.get("condicao")), None)
        itens = []
        for e in lista:
            if e.get("condicao"):
                continue
            textos = [f"Anexo {it.split('/')[0]}, item {it.split('/')[1]}: {anexos['itens'].get(it, '')}" for it in e["itens"]] or [f"Anexo {e['anexo']} (sem item específico na base)"]
            itens.append(f"{e['cclass']} ({'zero' if e['reducao'] == 100 else str(e['reducao']) + '%'}): " + " | ".join(textos))
        if not lista and condicionado is None:
            itens = ["Nenhum (código fora dos anexos)" if len("".join(ch for ch in codigo if ch.isdigit())) in (8, 9) else "Sem NCM/NBS válido para conferir"]
        livres = [e for e in lista if not e.get("condicao")]
        if len(livres) > 1:
            sugestao = "Depende do produto: " + " ou ".join(f"{e['cclass']} ({'zero' if e['reducao'] == 100 else str(e['reducao']) + '%'})" for e in livres)
        elif principal:
            sugestao = rotulo(principal)
        else:
            valido = len("".join(ch for ch in codigo if ch.isdigit())) == (9 if servico else 8)
            sugestao = "Só com condição (coluna K)" if condicionado else ("000001 — tributação integral" if valido else "Sem NCM/NBS válido: definir o código")
        leitura = cli.get("leituras", {}).get(codigo)
        resumo = leitura[0] if leitura else classificacao_auto(lista)
        valores = [
            fmt_codigo(codigo, servico), "Serviço" if servico else "Produto", v["descricao"], v["produtos"], v["valor"],
            f"=IF($E${ultima + 1}=0,0,E{r}/$E${ultima + 1})", ", ".join(v["cfops"]) or "—", v["pis_cofins"],
            "\n".join(itens) if itens else ("Só regra condicionada: " + rotulo(condicionado) if condicionado else ""),
            sugestao,
            f"{condicionado['cclass']} — {condicionado['condicao']}" if condicionado else "—",
            leitura[1] if leitura else "", resumo, "Pendente", "",
        ]
        fill = dict(CATEGORIAS).get(resumo)
        for c, val in enumerate(valores, start=1):
            cel = ws.cell(row=r, column=c, value=val)
            cel.font = font(color=RED if c == 11 and val != "—" else "222222", bold=c in (1, 13))
            cel.border = BORDER
            cel.alignment = Alignment(vertical="top", wrap_text=c in (3, 7, 9, 10, 11, 12, 13, 15))
            if c in (14, 15):
                cel.fill = PatternFill("solid", fgColor=INPUT)
            if c == 13 and fill:
                cel.fill = PatternFill("solid", fgColor=fill)
        ws.cell(row=r, column=1).number_format = "@"
        ws.cell(row=r, column=5).number_format = "#,##0.00"
        ws.cell(row=r, column=6).number_format = "0.0%"
    tot = ultima + 1
    ws.cell(row=tot, column=1, value="Total").font = font(bold=True)
    ws.cell(row=tot, column=4, value=f"=SUM(D{primeira}:D{ultima})").font = font(bold=True)
    ws.cell(row=tot, column=5, value=f"=SUM(E{primeira}:E{ultima})").font = font(bold=True)
    ws.cell(row=tot, column=6, value=f"=SUM(F{primeira}:F{ultima})").font = font(bold=True)
    ws.cell(row=tot, column=5).number_format = "#,##0.00"
    ws.cell(row=tot, column=6).number_format = "0.0%"
    for c in range(1, 16):
        ws.cell(row=tot, column=c).border = BORDER
        ws.cell(row=tot, column=c).fill = PatternFill("solid", fgColor="EEF1F4")
    dv = DataValidation(type="list", formula1='"Pendente,Revisar,Aprovado"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"N{primeira}:N{ultima}")
    ws.freeze_panes = ws.cell(row=primeira, column=2)
    ws.auto_filter.ref = f"A{head}:O{ultima}"
    return {"aba": cli["aba"], "primeira": primeira, "ultima": ultima, "periodo": periodo, "nome": cli["nome"], "cnpj": cnpj}


def resumo(wb, abas):
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = "NCMs e serviços vendidos pelos clientes — enquadramento na LC 214"
    ws["A1"].font = font(size=14, bold=True, color=NAVY)
    notas = [
        "De onde vem: resumo das notas fiscais de venda da Domínio, enviado pela Central (só leitura). Cada linha é um NCM (produto) ou NBS (serviço), com um produto de exemplo e quantos produtos diferentes foram vendidos nele.",
        "A “Leitura inicial” é uma primeira análise para o escritório conferir — não é parecer. A classificação final é por produto, olhando a descrição de cada um e o comprador.",
        "Para incluir outro cliente: cadastrar a empresa (nome e CNPJ) no módulo Reforma do AthosBI; na próxima rodada a Central manda o resumo fiscal dela. A lista completa de produtos vem da exportação do ERP do cliente.",
    ]
    for i, t in enumerate(notas, start=2):
        cel = ws.cell(row=i, column=1, value=t)
        cel.font = font(size=9, color=MUTED, italic=True)
        cel.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=10)
        ws.row_dimensions[i].height = 26
    head = 6
    colunas = [("Cliente", 34), ("CNPJ", 18), ("Período", 26), ("Vendas no período (R$)", 17), ("Códigos", 9)] + [(nome, 17) for nome, _ in CATEGORIAS] + [("% com redução ou zero", 12)]
    cabecalho(ws, head, colunas)
    for i, a in enumerate(abas):
        r = head + 1 + i
        q = f"'{a['aba']}'!"
        cats = f"{q}$M${a['primeira']}:$M${a['ultima']}"
        vals = f"{q}$E${a['primeira']}:$E${a['ultima']}"
        cnpj = a["cnpj"]
        linha = [a["nome"], f"{cnpj[:2]}.{cnpj[2:5]}.{cnpj[5:8]}/{cnpj[8:12]}-{cnpj[12:]}", a["periodo"], f"=SUM({vals})", f"=COUNTA({q}$A${a['primeira']}:$A${a['ultima']})"]
        linha += [f'=SUMIF({cats},"{nome}",{vals})' for nome, _ in CATEGORIAS]
        linha.append(f"=IF(D{r}=0,0,(F{r}+G{r})/D{r})")
        for c, val in enumerate(linha, start=1):
            cel = ws.cell(row=r, column=c, value=val)
            cel.font = font(bold=c == 1)
            cel.border = BORDER
            cel.alignment = Alignment(vertical="top", wrap_text=c in (1, 3))
            if 4 <= c <= 9 and c != 5:
                cel.number_format = "#,##0.00"
        ws.cell(row=r, column=10).number_format = "0.0%"
    for c, (nome, cor) in enumerate(CATEGORIAS, start=6):
        ws.cell(row=head, column=c).fill = PatternFill("solid", fgColor=NAVY)
        ws.cell(row=head + len(abas) + 2, column=c, value=nome).fill = PatternFill("solid", fgColor=cor)
        ws.cell(row=head + len(abas) + 2, column=c).font = font(size=9)
        ws.cell(row=head + len(abas) + 2, column=c).alignment = Alignment(wrap_text=True)
    ws.cell(row=head + len(abas) + 2, column=5, value="Legenda:").font = font(size=9, bold=True)
    ws.row_dimensions[head + len(abas) + 2].height = 26


def main():
    entrada, saida = Path(sys.argv[1]), Path(sys.argv[2])
    dados = json.loads(entrada.read_text(encoding="utf-8"))
    anexos = json.loads((DADOS / "anexos-lc214.json").read_text(encoding="utf-8"))
    wb = Workbook()
    abas = [aba_cliente(wb, anexos, cli, dados["sincronizado"]) for cli in dados["clientes"]]
    resumo(wb, abas)  # a aba padrão (a primeira) vira o Resumo
    wb.properties.creator = "Nota Certa"
    wb.save(saida)
    print("ok", saida, [a["aba"] for a in abas])


if __name__ == "__main__":
    main()
