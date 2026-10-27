"""Planilha base de trabalho do Nota Certa (3 abas):

- Referência LC 214: NCM/NBS dos anexos da LC 214 → item, cClassTrib, redução, condição;
- CFOP: os mais usados, tipo de operação e orientação inicial de IBS/CBS;
- Modelo de catálogo: uma linha por produto/serviço, com a sugestão da LC 214
  calculada por fórmula a partir da aba de referência.

Uso: python3 gerar-planilha.py "../planilhas/Nota Certa - Base de trabalho.xlsx"
Depois, recalcular as fórmulas (ex.: abrir e salvar no Excel/LibreOffice).
"""

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DADOS = Path(__file__).resolve().parent.parent / "dados"

NAVY = "14202B"
ORANGE = "D9622B"
MUTED = "5B6670"
INPUT = "FFF4D6"  # células para preencher
CALC = "EEF1F4"  # células calculadas
RED = "B42318"
FONT = "Arial"

thin = Side(style="thin", color="D5DBE0")
BORDER = Border(top=thin, bottom=thin, left=thin, right=thin)


def font(**kw):
    return Font(name=FONT, size=kw.pop("size", 10), **kw)


def titulo(ws, texto, subtitulos, ultima_coluna):
    ws["A1"] = texto
    ws["A1"].font = font(size=14, bold=True, color=NAVY)
    for i, linha in enumerate(subtitulos, start=2):
        ws.cell(row=i, column=1, value=linha).font = font(size=9, color=MUTED, italic=True)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=ultima_coluna)
        ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="top")


def cabecalho(ws, linha, colunas, fills=None):
    for c, (nome, largura) in enumerate(colunas, start=1):
        cel = ws.cell(row=linha, column=c, value=nome)
        cel.font = font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=(fills or {}).get(c, NAVY))
        cel.alignment = Alignment(wrap_text=True, vertical="center")
        cel.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = largura
    ws.row_dimensions[linha].height = 30


def formata_ncm(digitos):
    # 1006 · 1006.3 · 1006.30 · 1006.30.2 · 1006.30.21
    if len(digitos) <= 4:
        return digitos
    partes = [digitos[:4], digitos[4:6], digitos[6:8]]
    return ".".join(p for p in partes if p)


def formata_nbs(digitos):
    # 1 · 1.2201 · 1.2201.20 · 1.2201.20.00
    partes = [digitos[:1], digitos[1:5], digitos[5:7], digitos[7:9]]
    return ".".join(p for p in partes if p)


def referencia(wb, anexos):
    ws = wb.active
    ws.title = "Referência LC 214"
    colunas = [
        ("Tipo", 10), ("Código", 13), ("Abrange", 30), ("Exceções (não vale)", 18), ("Anexo", 7), ("Item", 6),
        ("Descrição do item no anexo", 48), ("Tratamento", 30), ("cClassTrib", 10), ("CST", 6), ("Redução", 9),
        ("Condição para valer", 34), ("Base legal", 14), ("Chave", 12),
    ]
    fonte = anexos["fonte"]
    titulo(ws, "Referência da LC 214 por NCM e NBS — reduções e alíquota zero", [
        "Como ler: o “Código” pode ser só o começo do NCM/NBS — a regra vale para todos os códigos que começam com ele, menos os da coluna “Exceções”. "
        "Um mesmo código pode aparecer em mais de uma linha: quem decide é a descrição do produto (coluna G) e, às vezes, o comprador (coluna L).",
        "NCM/NBS que não aparece aqui: alíquota cheia (cClassTrib 000001), salvo regimes específicos (combustíveis, financeiro, imóveis etc.). "
        "Medicamentos dependem também do registro na Anvisa (conferir).",
        f"Fonte: {fonte['pacote']} (dados {fonte['dados']}, licença Apache-2.0), extraído da Calculadora da Receita (módulo offline V0059) "
        "e do Informe Técnico 2025.002 v1.60. Material de estudo — confirmar na lei e nas tabelas oficiais vigentes.",
    ], len(colunas))
    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 26
    ws.row_dimensions[4].height = 26
    inicio = 6
    cabecalho(ws, inicio - 1, colunas)

    linhas = []
    for tipo, lista in (("Produto", anexos["ncm"]), ("Serviço", anexos["nbs"])):
        for prefixo, cclass, item, excecoes in lista:
            linhas.append((tipo, prefixo, cclass, item, excecoes))
    # Mesmo código: primeiro o que vale para qualquer comprador, depois a maior redução.
    trat = anexos["tratamentos"]
    linhas.sort(key=lambda r: (r[0] != "Produto", r[1], bool(trat[r[2]].get("condicao")), -trat[r[2]]["reducao"], r[3]))

    for i, (tipo, prefixo, cclass, item, excecoes) in enumerate(linhas):
        row = inicio + i
        t = trat[cclass]
        produto = tipo == "Produto"
        fmt = formata_ncm if produto else formata_nbs
        completo = len(prefixo) == (8 if produto else 9)
        abrange = ("Só este NCM" if produto else "Só esta NBS") if completo else (
            f"Todos os NCM que começam com {fmt(prefixo)}" if produto else f"Todas as NBS que começam com {fmt(prefixo)}")
        anexo, _, numero = (item or "/").partition("/")
        valores = [
            tipo, fmt(prefixo), abrange, ", ".join(fmt(e) for e in excecoes), anexo or t["anexo"], numero,
            anexos["itens"].get(item, ""), f"{t['titulo']} — {'alíquota zero' if t['reducao'] == 100 else 'redução de %d%%' % t['reducao']}",
            cclass, cclass[:3], t["reducao"] / 100, t.get("condicao", ""), t["artigo"], ("P" if produto else "S") + prefixo,
        ]
        for c, v in enumerate(valores, start=1):
            cel = ws.cell(row=row, column=c, value=v)
            cel.font = font(color=RED if c == 12 and v else "222222")
            cel.border = BORDER
            cel.alignment = Alignment(vertical="top", wrap_text=c in (3, 4, 7, 8, 12))
            if c in (2, 9, 10, 14):
                cel.number_format = "@"
        ws.cell(row=row, column=11).number_format = "0%"
    fim = inicio + len(linhas) - 1
    ws.freeze_panes = ws.cell(row=inicio, column=3)
    ws.auto_filter.ref = f"A{inicio - 1}:N{fim}"
    ws.column_dimensions["N"].hidden = True  # chave usada pelas fórmulas do catálogo
    ws.cell(row=inicio - 1, column=14).comment = Comment("Coluna auxiliar (P/S + dígitos) usada pelas fórmulas da aba Modelo de catálogo. Não apagar.", "Nota Certa")
    return inicio, fim


def cfop(wb, dados):
    ws = wb.create_sheet("CFOP")
    colunas = [("CFOP", 8), ("Sentido", 16), ("Equivalentes", 16), ("Descrição", 62), ("Tipo de operação", 24), ("IBS/CBS — orientação inicial", 62)]
    titulo(ws, "CFOP mais usados — tipo de operação e orientação para IBS/CBS", [
        dados["observacao"] + " O CFOP é da OPERAÇÃO, não do produto: o mesmo produto sai com CFOPs diferentes conforme a venda, a devolução, a remessa…",
        dados["fonte"],
    ], len(colunas))
    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 38
    inicio = 5
    cabecalho(ws, inicio - 1, colunas)
    sentido = {"1": "Entrada (no estado)", "2": "Entrada (outro estado)", "3": "Entrada (exterior)",
               "5": "Saída (no estado)", "6": "Saída (outro estado)", "7": "Saída (exterior)"}
    for i, item in enumerate(dados["cfops"]):
        row = inicio + i
        valores = [item["cfop"], sentido[item["cfop"][0]], item["equivalentes"], item["descricao"], item["tipo"], item["ibsCbs"]]
        for c, v in enumerate(valores, start=1):
            cel = ws.cell(row=row, column=c, value=v)
            cel.font = font(bold=c == 1)
            cel.border = BORDER
            cel.alignment = Alignment(vertical="top", wrap_text=c in (4, 5, 6))
            if c == 1:
                cel.number_format = "@"
    fim = inicio + len(dados["cfops"]) - 1
    ws.freeze_panes = ws.cell(row=inicio, column=2)
    ws.auto_filter.ref = f"A{inicio - 1}:F{fim}"


def catalogo(wb, ref_inicio, ref_fim, tratamentos):
    ws = wb.create_sheet("Modelo de catálogo")
    R = "'Referência LC 214'!"
    chave = f"{R}$N${ref_inicio}:$N${ref_fim}"

    colunas = [
        ("Código interno", 14), ("Descrição do produto ou serviço", 36), ("Tipo", 10), ("NCM (produto) ou NBS (serviço)", 16),
        ("GTIN (código de barras)", 16), ("Unidade", 8),
        ("Regras na LC 214", 10), ("Sugestão (cClassTrib)", 12), ("Tratamento sugerido", 34), ("Redução", 9),
        ("Item do anexo", 44), ("Condição", 30), ("Atenção", 34),
        ("Vende para órgão público ou entidade SUS?", 14), ("cClassTrib escolhido", 12), ("Status", 11), ("Aprovado por", 14),
        ("Data", 11), ("Observação", 30), ("Chave", 10), ("Linha da regra", 8),
    ]
    entrada = {1, 2, 3, 4, 5, 6, 14, 15, 16, 17, 18, 19}
    titulo(ws, "Modelo de catálogo — um produto ou serviço por linha", [
        "Preencha as colunas AMARELAS (A a F e N a S). As colunas CINZAS (G a M) são calculadas pela aba “Referência LC 214”: mostram a regra mais específica "
        "que o código encontrou. Quando há mais de uma regra (coluna G maior que 1) ou uma condição (coluna L), confira o item do anexo antes de decidir.",
        "A sugestão não considera as exceções dos anexos nem a descrição do produto: a decisão final (colunas O e P) é do escritório. "
        "Digite o NCM com 8 dígitos e a NBS com 9 (pode ser com ou sem pontos). As linhas 7 a 10 são exemplos — pode apagar.",
    ], 19)
    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 26
    head = 6
    fills = {c: ("8A5A00" if c in entrada else NAVY) for c in range(1, len(colunas) + 1)}
    cabecalho(ws, head, colunas, fills)
    ws.row_dimensions[head].height = 44

    exemplos = [
        ["ARR-005", "Arroz branco tipo 1, pacote 5 kg", "Produto", "1006.30.21", "7890000000017", "PCT", "Não", "200003", "Aprovado", "Escritório", None, "Exemplo — pode apagar"],
        ["SER-010", "Seringa descartável 10 ml com agulha", "Produto", "9018.31.19", "", "UN", "Sim", "", "Revisar", "", None, "Exemplo: definir 200030 (geral) e 200005 (órgão público/SUS)"],
        ["NB-001", "Notebook 14 polegadas", "Produto", "8471.30.12", "", "UN", "Não", "000001", "Aprovado", "Escritório", None, "Exemplo — pode apagar"],
        ["SV-EDU", "Mensalidade de ensino infantil", "Serviço", "1.2201.20.00", "", "MES", "Não", "", "Pendente", "", None, "Exemplo — pode apagar"],
    ]
    primeira = head + 1
    total = 300
    for i in range(total):
        r = primeira + i
        ex = exemplos[i] if i < len(exemplos) else None
        if ex:
            for c, v in zip((1, 2, 3, 4, 5, 6, 14, 15, 16, 17, 18, 19), ex):
                ws.cell(row=r, column=c, value=v)
        k = f"T{r}"
        ws[k] = (f'=IF(D{r}="","",IF(C{r}="Serviço","S","P")&SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(D{r},".",""),"-","")," ",""))')
        # Regra mais específica: o maior começo do código que existe na referência.
        busca = '""'
        for n in range(3, 11):
            busca = f"IFERROR(MATCH(LEFT({k},{n}),{chave},0),{busca})"
        ws[f"U{r}"] = f'=IF({k}="","",{busca})'
        ws[f"G{r}"] = f'=IF({k}="","",SUMPRODUCT(--(LEFT({k},LEN({chave}))={chave})))'
        linha = f"U{r}"
        idx = lambda col: f"INDEX({R}${col}${ref_inicio}:${col}${ref_fim},{linha})"  # MATCH já devolve a posição no intervalo
        ws[f"H{r}"] = f'=IF({k}="","",IF({linha}="","000001",{idx("I")}))'
        ws[f"I{r}"] = f'=IF({k}="","",IF({linha}="","Sem anexo: tributação integral",{idx("H")}))'
        ws[f"J{r}"] = f'=IF({k}="","",IF({linha}="",0,{idx("K")}))'
        ws[f"K{r}"] = f'=IF(OR({k}="",{linha}=""),"","Anexo "&{idx("E")}&", item "&{idx("F")}&": "&{idx("G")})'
        especial = f'SUMPRODUCT((LEFT({k},LEN({chave}))={chave})*({R}$L${ref_inicio}:$L${ref_fim}<>""))'
        ws[f"L{r}"] = (
            f'=IF(OR({k}="",{linha}=""),"",IF({idx("L")}&""<>"",{idx("L")}&"",'
            f'IF({especial}>0,"Há outro tratamento só para alguns compradores (ex.: órgão público/SUS) — veja a referência","")))'
        )
        digitos = f"LEN({k})-1"
        ws[f"M{r}"] = (
            f'=IF({k}="","",IF(AND(C{r}<>"Serviço",{digitos}<>8),"NCM deve ter 8 dígitos",'
            f'IF(AND(C{r}="Serviço",{digitos}<>9),"NBS deve ter 9 dígitos",'
            f'IF(L{r}<>"","Benefício depende do comprador: confira a coluna N",'
            f'IF(G{r}>1,"Mais de uma regra: confira o item do anexo na referência",'
            f'IF(G{r}=0,"Sem redução na LC 214","Confira se o produto é o descrito no item"))))))'
        )
        for c in range(1, len(colunas) + 1):
            cel = ws.cell(row=r, column=c)
            cel.border = BORDER
            cel.font = font(italic=bool(ex), color="222222" if c != 13 else "8A5A00")
            cel.alignment = Alignment(vertical="top", wrap_text=c in (2, 9, 11, 12, 13, 19))
            if c in entrada:
                cel.fill = PatternFill("solid", fgColor=INPUT)
            elif c <= 13:
                cel.fill = PatternFill("solid", fgColor=CALC)
            if c in (1, 4, 5, 8, 15):
                cel.number_format = "@"
        ws.cell(row=r, column=10).number_format = "0%"
        ws.cell(row=r, column=18).number_format = "DD/MM/YYYY"
    ultima = primeira + total - 1

    listas = [
        (f"C{primeira}:C{ultima}", '"Produto,Serviço"'),
        (f"N{primeira}:N{ultima}", '"Sim,Não"'),
        (f"P{primeira}:P{ultima}", '"Pendente,Revisar,Aprovado"'),
        (f"O{primeira}:O{ultima}", '"' + ",".join(["000001", *sorted(tratamentos)]) + '"'),
    ]
    for intervalo, formula in listas:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        if intervalo.startswith("O"):
            dv.showErrorMessage = False  # aceita outros códigos da tabela oficial
        ws.add_data_validation(dv)
        dv.add(intervalo)

    ws.column_dimensions["T"].hidden = True
    ws.column_dimensions["U"].hidden = True
    ws.freeze_panes = ws.cell(row=primeira, column=3)
    ws.auto_filter.ref = f"A{head}:S{ultima}"
    ws.cell(row=head, column=15).comment = Comment(
        "A classificação que vai para a nota. Se o produto tem regra diferente para órgão público/SUS, "
        "registre as duas (ex.: 200030 / 200005) e explique na observação.", "Nota Certa")


def main():
    saida = Path(sys.argv[1] if len(sys.argv) > 1 else "Nota Certa - Base de trabalho.xlsx")
    anexos = json.loads((DADOS / "anexos-lc214.json").read_text(encoding="utf-8"))
    cfops = json.loads((DADOS / "cfop-mais-usados.json").read_text(encoding="utf-8"))
    wb = Workbook()
    ref_inicio, ref_fim = referencia(wb, anexos)
    cfop(wb, cfops)
    catalogo(wb, ref_inicio, ref_fim, anexos["tratamentos"])
    wb.properties.creator = "Nota Certa"
    wb.properties.title = "Nota Certa — Base de trabalho"
    saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(saida)
    print("ok", saida, ref_fim - ref_inicio + 1, "regras", len(cfops["cfops"]), "CFOPs")


if __name__ == "__main__":
    main()
